from langchain_core.documents import Document
# Data => Documents
import os
import json
#from langchain_community.document_loaders.pdf import PyPDFLoader
from pypdf import PdfReader

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter



def load_regulatory_pdfs(folder_path="data/regulatory"):
  
    total_pdfs = 0
    all_docs = []

    if not os.path.exists(folder_path):
        print(f"Directory not found: {folder_path}")
        return all_docs

    for filename in os.listdir(folder_path):
        if filename.lower().endswith(".pdf"):
            pdf_path = os.path.join(folder_path, filename)
            total_pdfs += 1

            loader = PdfReader(pdf_path)
            for page_num, page in enumerate(loader.pages):
                text = page.extract_text() or ""

                doc = Document(
                    page_content=text,
                    metadata={"source": filename,"type": "regulatory_pdf","page": page_num+1}
                )
                all_docs.append(doc)


    print("Total PDFs processed:", total_pdfs)
    print("total pages loaded :", len(all_docs))
    return all_docs


def load_enterprise_policies(folder_path="data/policies"):
    all_docs = []
    total_policies = 0

    if not os.path.exists(folder_path):
        print(f"Directory not found: {folder_path}")
        return all_docs

    for filename in os.listdir(folder_path):
        if filename.lower().endswith(".docx"):
            docx_path = os.path.join(folder_path, filename)
            total_policies += 1

            doc_reader = DocxDocument(docx_path)
            for para_num, paragraph in enumerate(doc_reader.paragraphs):
                text = paragraph.text.strip()
                if text:
                    doc = Document(
                        page_content=text,
                        metadata={"source": filename, "type": "enterprise_policy", "paragraph": para_num + 1},
                        
                    )
                    all_docs.append(doc)

    print(f"Total Enterprise Policies processed: {total_policies}")
    return all_docs

    
#chunks

def split_doc(documents, chunk_size=500, chunk_overlap=50):

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size = chunk_size,
        chunk_overlap = chunk_overlap,
        separators=["\n\n", "\n", " ", ""]
    )

    chunked_docs = text_splitter.split_documents(documents)
    return chunked_docs

def save_chunks_to_json(chunks, output_path="data/corpus/parsed_corpus_chunks.json"):
    """Saving split documents into a structured JSON corpus """
    #os.path.dirname(output_path)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    serializable_chunks = [
        {
            "chunk_id": idx,
            "page_content": chunk.page_content,
            "metadata": chunk.metadata
        }
        for idx, chunk in enumerate(chunks)
    ]
    
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(serializable_chunks, f, indent=4)
    print(f"Saved {len(serializable_chunks)} structured chunks to JSON corpus: {output_path}")


if __name__ == "__main__":
    print("Running complete Stage 1 Ingestion Pipeline...")
    
    # Load both regulatory documents and enterprise policies
    regulatory_docs = load_regulatory_pdfs("data/regulatory")
    enterprise_docs = load_enterprise_policies("data/policies")
    
    # Combine them into a single raw documents list
    raw_documents = regulatory_docs + enterprise_docs
    
    print(f"Total raw documents loaded: {len(raw_documents)}")
    
    if raw_documents:
        chunks = split_doc(raw_documents)
        print("Success! Total chunks created:", len(chunks))
        
        # Save to JSON corpus
        save_chunks_to_json(chunks)
        
        print("-" * 40)
        print("First chunk sample:")
        print(chunks[0].page_content[:300])
        print("Metadata:", chunks[0].metadata)